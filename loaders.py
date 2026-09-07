"""Drive/xlsx -> normalized frames.

Every client sheet is shaped differently, so each gets a small adapter that
emits the SAME two frames:

  budgets   : client | account | allocated_budget | target_cpa | status | days_running
  campaigns : date | account | campaign | campaign_type | impressions | clicks
              | cost | conversions | conv_value

Add a new client by writing one adapter function. Nothing else changes.
"""
from __future__ import annotations

import base64
import json
import re
from pathlib import Path

import openpyxl
import pandas as pd

RAW = Path(__file__).parent / "raw"

CAMPAIGN_COLS = ["date", "account", "campaign", "campaign_type",
                 "impressions", "clicks", "cost", "conversions", "conv_value"]


# --------------------------------------------------------------------------- #
def save_drive_payload(tool_result_path: str, out_name: str) -> Path:
    """Decode a Drive download_file_content payload into raw/<out_name>.xlsx."""
    d = json.load(open(tool_result_path))
    RAW.mkdir(exist_ok=True)
    p = RAW / f"{out_name}.xlsx"
    p.write_bytes(base64.b64decode(d["content"]))
    return p


def norm(s) -> str:
    """Loose key for matching 'Cary Family' <-> 'CaryFamily' <-> 'cary family'."""
    return re.sub(r"[^a-z0-9]", "", str(s or "").lower())


# --------------------------------------------------------------------------- #
def load_budgets(path: Path | str = None) -> pd.DataFrame:
    """Budget Tracker -> the budget frame. This is the source of truth."""
    path = Path(path or RAW / "budget_tracker.xlsx")
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["DPR"] if "DPR" in wb.sheetnames else wb.worksheets[-1]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(c).strip() if c else "" for c in rows[0]]
    df = pd.DataFrame(rows[1:], columns=hdr)
    df = df.rename(columns={
        "Client": "client", "Account": "account",
        "Allocated Budget": "allocated_budget",
        "Target CPA/ROAS/CPM": "target_cpa",
        "Account Status": "status",
        "Days Ads are Running": "days_running",
        "AM": "am", "Channels": "channel", "Date": "budget_month",
    })
    df = df[df["client"].notna() & df["account"].notna()].copy()
    df["allocated_budget"] = pd.to_numeric(
        df["allocated_budget"].astype(str).str.replace(r"[,$\s]", "", regex=True),
        errors="coerce").fillna(0.0)
    df["target_cpa"] = pd.to_numeric(df["target_cpa"], errors="coerce")
    for c in ("client", "account", "status", "days_running", "am", "channel"):
        if c in df:
            df[c] = df[c].astype(str).str.strip()
    df["key"] = df["account"].map(norm)
    return df.reset_index(drop=True)


# --------------------------------------------------------------------------- #
# Adapters — one per client sheet layout
# --------------------------------------------------------------------------- #
def adapter_supermetrics_data_tab(path: Path | str, account_col: str | None,
                                  fallback_account: str | None = None,
                                  lookup_sheet: str | None = "Campaign Lookup",
                                  ) -> pd.DataFrame:
    """The standard Position2 layout: a `Data` tab written by Supermetrics.

    Columns: Date | Campaign name | Impressions | Clicks | Cost | Conversions
             | (All|Total) conversion value | [Location] | Month | Campaign Type

    account_col      : column holding the sub-account/location ('Location'), or
                       None for single-account clients.
    fallback_account : account name to stamp on every row when account_col is None.
    lookup_sheet     : optional Campaign -> Location/Type map used to repair
                       missing values (Supermetrics lookups break constantly).
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Data"]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(c).strip() if c else "" for c in rows[0]]
    df = pd.DataFrame(rows[1:], columns=hdr)
    df = df[df[hdr[0]].notna()]

    def pick(*names):
        for n in names:
            if n in df.columns:
                return n
        return None

    out = pd.DataFrame({
        "date": pd.to_datetime(df[pick("Date")], errors="coerce"),
        "campaign": df[pick("Campaign name", "Campaign")].astype(str).str.strip(),
        "impressions": pd.to_numeric(df[pick("Impressions")], errors="coerce").fillna(0),
        "clicks": pd.to_numeric(df[pick("Clicks")], errors="coerce").fillna(0),
        "cost": pd.to_numeric(df[pick("Cost")], errors="coerce").fillna(0),
        "conversions": pd.to_numeric(df[pick("Conversions")], errors="coerce").fillna(0),
        "conv_value": pd.to_numeric(
            df[pick("All conversion value", "Total conversion value")],
            errors="coerce").fillna(0),
    })
    ctype = pick("Campaign Type")
    out["campaign_type"] = df[ctype].astype(str).str.strip() if ctype else "Unknown"

    if account_col and account_col in df.columns:
        out["account"] = df[account_col].astype(str).str.strip()
    else:
        out["account"] = fallback_account or "Account"

    # --- repair via the Campaign Lookup tab -------------------------------
    if lookup_sheet and lookup_sheet in wb.sheetnames:
        lk = list(wb[lookup_sheet].iter_rows(values_only=True))
        lh = [str(c).strip() if c else "" for c in lk[0]]
        L = pd.DataFrame(lk[1:], columns=lh)
        L = L[L[lh[0]].notna()]
        cmap = {str(r[lh[0]]).strip(): r for _, r in L.iterrows()}
        loc_col = next((c for c in lh if c.lower() == "location"), None)
        typ_col = next((c for c in lh if "brand" in c.lower() or "type" in c.lower()), None)

        def fix_type(row):
            v = str(row["campaign_type"])
            if v and v.lower() not in ("none", "nan", "#n/a", "unknown", ""):
                return v
            r = cmap.get(row["campaign"])
            return str(r[typ_col]).strip() if r is not None and typ_col else "Unknown"

        def fix_acct(row):
            v = str(row["account"])
            if v and v.lower() not in ("none", "nan", "#n/a", ""):
                return v
            r = cmap.get(row["campaign"])
            return str(r[loc_col]).strip() if r is not None and loc_col else v

        out["campaign_type"] = out.apply(fix_type, axis=1)
        if account_col:
            out["account"] = out.apply(fix_acct, axis=1)

        # prefix-match fallback for names like P2_Brand_McH_Set-A
        unknown = out["campaign_type"].str.lower().isin(["unknown", "nan", "none", "#n/a"])
        if unknown.any() and typ_col:
            keys = sorted(cmap, key=len, reverse=True)
            def by_prefix(c):
                for k in keys:
                    if c.startswith(k):
                        return str(cmap[k][typ_col]).strip()
                return "Unknown"
            out.loc[unknown, "campaign_type"] = out.loc[unknown, "campaign"].map(by_prefix)

    out = out[out["date"].notna()]
    out["key"] = out["account"].map(norm)
    return out[CAMPAIGN_COLS + ["key"]].reset_index(drop=True)


# --------------------------------------------------------------------------- #
def month_slice(campaigns: pd.DataFrame, year: int, month: int) -> pd.DataFrame:
    m = (campaigns["date"].dt.year == year) & (campaigns["date"].dt.month == month)
    return campaigns[m].copy()


def reconcile(campaigns: pd.DataFrame, budgets: pd.DataFrame,
              client: str) -> tuple[pd.DataFrame, list[str]]:
    """Align campaign `account` values onto Budget Tracker account names.

    Returns the aligned frame plus a list of human-readable warnings.
    """
    b = budgets[budgets["client"].str.strip().str.lower() == client.lower()]
    bk = dict(zip(b["key"], b["account"]))
    warns: list[str] = []

    campaigns = campaigns.copy()
    campaigns["account"] = campaigns["key"].map(lambda k: bk.get(k))
    unmatched = campaigns["account"].isna()
    if unmatched.any():
        lost = (campaigns[unmatched].groupby("key")["cost"].sum()
                .sort_values(ascending=False))
        for k, v in lost.items():
            if v > 0:
                warns.append(f"No budget row for '{k}' (${v:,.0f} MTD spend) — excluded.")
        campaigns = campaigns[~unmatched]

    have = set(campaigns["key"])
    for _, r in b.iterrows():
        if r["key"] not in have and r["allocated_budget"] > 0:
            warns.append(f"'{r['account']}' has a ${r['allocated_budget']:,.0f} budget "
                         f"but no campaign data.")
    return campaigns, warns
